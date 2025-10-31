import docx
import PyPDF2
import pdfplumber
from typing import Dict, Optional
from pathlib import Path
import re


class DocumentProcessor:
    """Process uploaded MOP documents (PDF, DOCX)."""
    
    def __init__(self):
        pass
    
    async def extract_text(self, file_path: str) -> str:
        """Extract text from document based on file type."""
        path = Path(file_path)
        extension = path.suffix.lower()
        
        if extension == '.pdf':
            return await self._extract_from_pdf(file_path)
        elif extension in ['.docx', '.doc']:
            return await self._extract_from_docx(file_path)
        else:
            raise ValueError(f"Unsupported file format: {extension}")
    
    async def _extract_from_pdf(self, file_path: str) -> str:
        """Extract text from PDF file."""
        text = ""
        try:
            # Try pdfplumber first (better formatting)
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
        except Exception:
            # Fallback to PyPDF2
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    text += page.extract_text() + "\n"
        
        return text.strip()
    
    async def _extract_from_docx(self, file_path: str) -> str:
        """Extract text from DOCX file."""
        doc = docx.Document(file_path)
        text = []
        
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text.append(paragraph.text)
        
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text.append(cell.text)
        
        return "\n".join(text)
    
    async def extract_sections(self, text: str) -> Dict[str, str]:
        """
        Extract MOP sections: Pre-Checks, Operation Steps, Rollback Steps.
        Uses pattern matching and heuristics.
        """
        sections = {
            "pre_checks": "",
            "operation_steps": "",
            "rollback_steps": "",
            "full_text": text
        }
        
        # Normalize text
        text_normalized = text.lower()
        
        # Define section patterns
        pre_check_patterns = [
            r'pre[\s-]*checks?',
            r'prerequisites?',
            r'pre[\s-]*requisites?',
            r'before[\s-]*you[\s-]*begin',
            r'preparation',
            r'pre[\s-]*conditions?'
        ]
        
        operation_patterns = [
            r'operation[\s-]*steps?',
            r'implementation[\s-]*steps?',
            r'procedure[\s-]*steps?',
            r'execution[\s-]*steps?',
            r'main[\s-]*steps?',
            r'change[\s-]*steps?'
        ]
        
        rollback_patterns = [
            r'rollback[\s-]*steps?',
            r'rollback[\s-]*procedure',
            r'backout[\s-]*plan',
            r'recovery[\s-]*steps?',
            r'reversal[\s-]*procedure'
        ]
        
        # Find section boundaries
        section_positions = []
        
        for pattern in pre_check_patterns:
            matches = list(re.finditer(pattern, text_normalized))
            for match in matches:
                section_positions.append(('pre_checks', match.start()))
        
        for pattern in operation_patterns:
            matches = list(re.finditer(pattern, text_normalized))
            for match in matches:
                section_positions.append(('operation_steps', match.start()))
        
        for pattern in rollback_patterns:
            matches = list(re.finditer(pattern, text_normalized))
            for match in matches:
                section_positions.append(('rollback_steps', match.start()))
        
        # Sort by position
        section_positions.sort(key=lambda x: x[1])
        
        # Extract text between sections
        for i, (section_name, start_pos) in enumerate(section_positions):
            end_pos = section_positions[i + 1][1] if i + 1 < len(section_positions) else len(text)
            section_text = text[start_pos:end_pos].strip()
            
            # Only keep the first occurrence of each section
            if not sections[section_name]:
                sections[section_name] = section_text
        
        return sections
