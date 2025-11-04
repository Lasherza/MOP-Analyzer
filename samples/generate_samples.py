#!/usr/bin/env python3
"""
Generate sample DOCX MOP files for testing
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import os

def create_full_mop():
    """Create a complete, well-structured MOP"""
    doc = Document()
    
    # Title
    title = doc.add_heading('Network Configuration Change - IPCORE', 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    
    # Metadata
    doc.add_paragraph('Change ID: CHG0001234')
    doc.add_paragraph('Category: IPCORE')
    doc.add_paragraph('Owner: John Doe')
    doc.add_paragraph('Date: 2024-01-15')
    doc.add_paragraph('')
    
    # Pre-Checks
    doc.add_heading('Pre-Checks', 1)
    pre_checks = [
        'Verify change approval received from CAB (Change Advisory Board)',
        'Confirm backup of current configuration completed successfully',
        'Check system health status - all services running normally',
        'Verify all dependencies are met and prerequisite changes deployed',
        'Risk assessment completed and mitigation plan approved by owner',
        'Ensure maintenance window scheduled and stakeholders notified',
        'Verify required resources (personnel, tools, access) are available',
        'Confirm rollback plan tested and validated'
    ]
    for i, check in enumerate(pre_checks, 1):
        doc.add_paragraph(f'{i}. {check}', style='List Number')
    doc.add_paragraph('')
    
    # Operation Steps
    doc.add_heading('Operation Steps', 1)
    doc.add_paragraph('Responsible: Network Operations Team')
    doc.add_paragraph('')
    
    op_steps = [
        'Notify all stakeholders that change is starting',
        'Take snapshot of current system state for validation',
        'Stop the target service gracefully',
        'Verify service has stopped completely - check process status',
        'Backup current configuration files to /backup/config_YYYYMMDD/',
        'Update configuration parameters as per change request CR-12345',
        'Validate configuration syntax using config-check tool',
        'Start the service and monitor startup logs',
        'Verify service is running - check process and port status',
        'Test connectivity to upstream and downstream systems',
        'Monitor system metrics for 15 minutes - CPU, memory, network',
        'Validate expected outcomes: response time < 100ms, error rate < 0.1%',
        'Update CMDB with new configuration version',
        'Notify stakeholders of successful completion'
    ]
    for i, step in enumerate(op_steps, 1):
        doc.add_paragraph(f'{i}. {step}', style='List Number')
    
    doc.add_paragraph('')
    doc.add_paragraph('Note: If any validation fails, proceed to rollback immediately.', style='Intense Quote')
    doc.add_paragraph('')
    
    # Rollback Steps
    doc.add_heading('Rollback Steps', 1)
    doc.add_paragraph('Rollback Trigger Criteria:')
    triggers = [
        'Service fails to start after 3 attempts',
        'Critical errors in application logs',
        'Performance degradation > 20%',
        'Failed connectivity tests',
        'Decision by Change Manager or Service Owner'
    ]
    for trigger in triggers:
        doc.add_paragraph(trigger, style='List Bullet')
    
    doc.add_paragraph('')
    doc.add_paragraph('Rollback Procedure (Owner: Change Manager):')
    rollback_steps = [
        'Stop the service immediately',
        'Restore original configuration from backup location',
        'Verify configuration file integrity using checksum',
        'Start the service with original configuration',
        'Validate system returns to baseline state',
        'Confirm all metrics return to normal levels',
        'Notify stakeholders of rollback completion',
        'Document rollback reason and lessons learned'
    ]
    for i, step in enumerate(rollback_steps, 1):
        doc.add_paragraph(f'{i}. {step}', style='List Number')
    
    doc.add_paragraph('')
    doc.add_paragraph('Contact: change.manager@telco.com | Emergency: +1-555-0100')
    
    # Save
    output_path = os.path.join(os.path.dirname(__file__), 'full_mop.docx')
    doc.save(output_path)
    print(f'Created: {output_path}')

def create_partial_mop():
    """Create a MOP with some missing elements"""
    doc = Document()
    
    doc.add_heading('Database Maintenance - Packet Core', 0)
    doc.add_paragraph('')
    
    # Pre-Checks (incomplete)
    doc.add_heading('Pre-Checks', 1)
    doc.add_paragraph('1. Check database backup', style='List Number')
    doc.add_paragraph('2. Verify system status', style='List Number')
    doc.add_paragraph('')
    
    # Operation Steps (missing verification and owner)
    doc.add_heading('Procedure', 1)
    doc.add_paragraph('1. Stop database service')
    doc.add_paragraph('2. Run maintenance script')
    doc.add_paragraph('3. Start database service')
    doc.add_paragraph('')
    
    # Rollback (incomplete - no triggers or owner)
    doc.add_heading('Rollback Plan', 1)
    doc.add_paragraph('1. Restore from backup if needed')
    doc.add_paragraph('2. Restart service')
    doc.add_paragraph('')
    
    output_path = os.path.join(os.path.dirname(__file__), 'partial_mop.docx')
    doc.save(output_path)
    print(f'Created: {output_path}')

def create_poor_mop():
    """Create a poorly structured MOP"""
    doc = Document()
    
    doc.add_heading('System Update', 0)
    doc.add_paragraph('')
    
    # Very minimal pre-checks
    doc.add_heading('Checks', 1)
    doc.add_paragraph('Make sure system is ready')
    doc.add_paragraph('')
    
    # Vague operation steps
    doc.add_heading('Steps', 1)
    doc.add_paragraph('1. Do the update')
    doc.add_paragraph('2. Check if it works')
    doc.add_paragraph('')
    
    # No real rollback
    doc.add_heading('If it fails', 1)
    doc.add_paragraph('Undo the changes')
    doc.add_paragraph('')
    
    output_path = os.path.join(os.path.dirname(__file__), 'poor_mop.docx')
    doc.save(output_path)
    print(f'Created: {output_path}')

if __name__ == '__main__':
    print('Generating sample MOP documents...')
    create_full_mop()
    create_partial_mop()
    create_poor_mop()
    print('Done!')
